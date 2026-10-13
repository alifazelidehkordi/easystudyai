#!/usr/bin/env python3
"""Convert professor PDFs, PowerPoints, Word files or text to cited Markdown.

Everything runs locally. This script never connects to any LLM/API.
OCR is intentionally not attempted: scanned pages require visual review.
"""
from __future__ import annotations

import argparse
import hashlib
import re
import sys
from pathlib import Path
from datetime import datetime, timezone


def slug(text: str) -> str:
    return (re.sub(r"[^a-zA-Z0-9_-]+", "-", text).strip("-")[:70] or "lecture").lower()


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def tidy(value: str) -> str:
    return (value or "").replace("\x00", "").replace("\r\n", "\n").strip()


def read_pdf(source: Path, assets: Path, render_pages: bool):
    try:
        import pymupdf
    except ImportError as exc:
        raise RuntimeError("PDF support requires pip install -r requirements.txt") from exc
    result, warnings = [], []
    with pymupdf.open(source) as pdf:
        if pdf.is_encrypted:
            raise RuntimeError("Encrypted PDF: unlock a permitted copy first")
        for number, page in enumerate(pdf, 1):
            text = tidy(page.get_text("text", sort=True))
            chunk = [f"## Page {number}", "", text or "[No selectable text: check scanned page image.]"]
            if not text:
                warnings.append(f"Page {number}: no text layer; OCR not performed")
            if render_pages:
                assets.mkdir(parents=True, exist_ok=True)
                filename = f"page-{number:04}.png"
                page.get_pixmap(dpi=120, alpha=False).save(str(assets / filename))
                chunk += ["", f"![Lecture PDF page {number}](assets/{assets.name}/{filename})"]
            result.append("\n".join(chunk) + "\n")
    return result, warnings


def read_pptx(source: Path, assets: Path):
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise RuntimeError("PowerPoint support requires pip install -r requirements.txt") from exc
    sections, warnings = [], []
    for n, slide in enumerate(Presentation(str(source)).slides, 1):
        lines = [f"## Slide {n}", ""]
        images = 0
        for shape in slide.shapes:
            if shape.has_text_frame and tidy(shape.text):
                lines.extend([tidy(shape.text), ""])
            if shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(tidy(cell.text).replace("\n", " ") for cell in row.cells))
                lines.append("")
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                ext = shape.image.ext.lower()
                if ext not in {"png", "jpg", "jpeg", "webp", "gif"}:
                    warnings.append(f"Slide {n}: skipped image type {ext}")
                    continue
                images += 1
                assets.mkdir(parents=True, exist_ok=True)
                filename = f"slide-{n:04}-figure-{images:02}.{ext}"
                (assets / filename).write_bytes(shape.image.blob)
                lines.extend([f"![Original figure, slide {n}](assets/{assets.name}/{filename})", ""])
        if len(lines) <= 2:
            lines.append("[No readable text/inline figure; inspect original slide.]")
        sections.append("\n".join(lines) + "\n")
    warnings.append("Slide layouts, groups and diagrams may not survive extraction: inspect original slides.")
    return sections, warnings


def read_docx(source: Path, assets: Path):
    try:
        from docx import Document
        from docx.parts.image import ImagePart
    except ImportError as exc:
        raise RuntimeError("Word support requires pip install -r requirements.txt") from exc
    document = Document(source)
    chunks = ["## Document", ""]
    for n, paragraph in enumerate(document.paragraphs, 1):
        text = tidy(paragraph.text)
        if text:
            chunks.extend([f"Paragraph {n}: {text}", ""])
    for n, table in enumerate(document.tables, 1):
        chunks.extend([f"### Table {n}", ""])
        for row in table.rows:
            chunks.append(" | ".join(tidy(cell.text).replace("\n", " ") for cell in row.cells))
        chunks.append("")
    # Images may be attached outside the paragraph sequence.
    seen = set()
    for part in document.part.related_parts.values():
        if not isinstance(part, ImagePart):
            continue
        ext = part.filename.rsplit(".", 1)[-1].lower()
        if ext not in {"png", "jpg", "jpeg", "webp", "gif"}:
            continue
        fingerprint = hashlib.sha256(part.blob).hexdigest()[:10]
        if fingerprint in seen:
            continue
        seen.add(fingerprint)
        assets.mkdir(parents=True, exist_ok=True)
        filename = f"doc-figure-{fingerprint}.{ext}"
        (assets / filename).write_bytes(part.blob)
        chunks.extend([f"![DOCX figure; position not established](assets/{assets.name}/{filename})", ""])
    return ["\n".join(chunks)], [
        "DOCX paragraph numbers are not printed page numbers; figure/table ordering may differ from original."
    ]


def convert(source: Path, destination: Path, render_pages: bool, overwrite: bool):
    if not source.is_file():
        raise ValueError(f"File not found: {source}")
    ext = source.suffix.lower()
    if ext not in {".pdf", ".pptx", ".docx", ".txt", ".md"}:
        raise ValueError(f"Unsupported file type: {ext}")
    destination.mkdir(parents=True, exist_ok=True)
    stem = slug(source.stem) + "-" + ext.lstrip(".")
    output = destination / (stem + ".md")
    if output.exists() and not overwrite:
        raise ValueError(f"Already imported: {output} (use --force to replace)")
    assets = destination / "assets" / stem
    if ext == ".pdf":
        sections, warnings = read_pdf(source, assets, render_pages)
    elif ext == ".pptx":
        sections, warnings = read_pptx(source, assets)
    elif ext == ".docx":
        sections, warnings = read_docx(source, assets)
    else:
        sections, warnings = ["## Source text\n\n" + tidy(source.read_text(encoding="utf-8-sig"))], []
    # Keep absolute paths, patient/student names, and environment details out of generated citations.
    header = [
        f"# Imported lecture: {source.name}",
        "",
        f"- Original filename: {source.name}",
        f"- Source SHA256: {sha256(source)}",
        f"- Imported at (UTC): {datetime.now(timezone.utc).isoformat(timespec='seconds')}",
        "- Provenance: lecturer/student-provided, NOT externally verified",
        "- Cite the original filename and the exact page/slide/paragraph anchors below",
        "- Inspect layouts, figure labels, table alignment and all scan-only pages",
        "",
    ]
    if warnings:
        header.extend(["## Import warnings", "", *["- " + warning for warning in warnings], ""])
    output.write_text("\n".join(header + sections), encoding="utf-8")
    return output, len(sections), warnings


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    parser.add_argument("--output", type=Path, default=Path("course/processed"))
    parser.add_argument("--render-pages", action="store_true", help="Render every PDF page as PNG")
    parser.add_argument("--force", action="store_true", help="Replace existing Markdown")
    args = parser.parse_args()
    try:
        output, n, warnings = convert(args.file.expanduser(), args.output, args.render_pages, args.force)
        print(f"Imported {n} section(s) to {output}; warnings: {len(warnings)}")
        for warning in warnings:
            print("CHECK:", warning)
        return 0
    except (OSError, ValueError, RuntimeError) as exc:
        print("ERROR:", exc, file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
