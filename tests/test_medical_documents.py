"""Medical-style documents exercising real PDF, PPTX, and DOCX import paths.

Fixtures are created in a temporary directory so nothing resembling professor
notes or clinical data is committed. No network or LLM provider is used.
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent


def cli(script: str, *args: object, expected: int = 0):
    r = subprocess.run(
        [sys.executable, str(ROOT / "scripts" / script), *map(str, args)],
        cwd=ROOT, capture_output=True, text=True, timeout=45
    )
    if r.returncode != expected:
        raise AssertionError(
            f"{script} returned {r.returncode} != {expected}\n{r.stdout}\n{r.stderr}"
        )
    return r


def png(path: Path) -> None:
    from PIL import Image
    # A diagram-like fixture, not an authentic medical specimen.
    image = Image.new("RGB", (180, 80), "white")
    image.save(path, format="PNG")


class MedicalDocumentTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.notes = self.root / "processed"
        self.briefs = self.root / "briefs"
        self.validation = self.root / "validation"

    def test_pdf_with_vector_diagram_and_unreadable_page(self):
        try:
            import pymupdf
            from PIL import Image
        except ImportError:
            self.skipTest("optional PDF/Pillow dependencies unavailable")
        input_file = self.root / "Cardiac anatomy.pdf"
        diagram = self.root / "figure.png"
        png(diagram)
        document = pymupdf.open()
        page1 = document.new_page()
        page1.insert_text((50, 60), "Cardiac cycle: systolic pressure and aortic outflow.")
        page1.draw_line((70, 90), (160, 140), color=(1, 0, 0), width=3)
        page2 = document.new_page()
        page2.insert_image(pymupdf.Rect(60, 60, 340, 260), filename=str(diagram))
        document.save(str(input_file))
        document.close()
        before = input_file.read_bytes()

        cli("learn.py", input_file, "--notes", self.notes, "--briefs", self.briefs)
        imported = self.notes / "cardiac-anatomy-pdf.md"
        self.assertTrue((self.notes / "assets/cardiac-anatomy-pdf/page-0001.png").is_file())
        self.assertTrue((self.notes / "assets/cardiac-anatomy-pdf/page-0002.png").is_file())
        self.assertIn("No selectable text", imported.read_text())
        self.assertEqual(before, input_file.read_bytes())

        check = cli(
            "validate_lecture.py", imported, "--original", input_file,
            "--output", self.validation, "--json", "--strict", expected=2
        )
        info = json.loads(check.stdout)
        self.assertEqual(info["original_validation"], "matching SHA256")
        self.assertEqual(info["original_page_or_slide_count"], 2)
        self.assertEqual(info["no_text_sections"], ["Page 2"])
        self.assertEqual(info["technical_status"], "needs_review")
        report = self.validation / "cardiac-anatomy-pdf-validation.md"
        self.assertIn("not proof of factual correctness", report.read_text())

        input_file.write_bytes(before + b"\n% revised source")
        stale = cli(
            "validate_lecture.py", imported, "--original", input_file,
            "--output", self.validation, "--json", "--strict", expected=2
        )
        self.assertTrue(any("differs" in x for x in json.loads(stale.stdout)["blocking_issues"]))

    def test_pptx_lecture_tables_notes_and_figures(self):
        try:
            from pptx import Presentation
            from pptx.util import Inches
            from PIL import Image
        except ImportError:
            self.skipTest("optional PowerPoint/Pillow unavailable")
        slide_path = self.root / "Physiology slides.pptx"
        picture = self.root / "schematic.png"
        png(picture)

        ppt = Presentation()
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(8), Inches(1)).text = (
            "Renal blood flow and glomerular filtration"
        )
        table = slide.shapes.add_table(2, 2, Inches(0.6), Inches(1.5), Inches(6), Inches(2)).table
        table.cell(0, 0).text = "Parameter"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Blood flow"
        table.cell(1, 1).text = "High in renal cortex"
        slide.shapes.add_picture(str(picture), Inches(2), Inches(4), width=Inches(3))
        notes = slide.notes_slide.notes_text_frame
        notes.text = "Speaker emphasis: compare afferent and efferent arterioles."
        ppt.save(str(slide_path))
        cli("learn.py", slide_path, "--notes", self.notes, "--briefs", self.briefs)
        imported = self.notes / "physiology-slides-pptx.md"
        body = imported.read_text(encoding="utf-8")
        self.assertIn("## Slide 1", body)
        self.assertIn("Parameter | Value", body)
        self.assertIn("Blood flow | High in renal cortex", body)
        self.assertIn("Speaker notes (slide 1;", body)
        self.assertIn("compare afferent and efferent", body)
        self.assertIn("![Original figure, slide 1]", body)
        self.assertTrue((self.notes / "assets/physiology-slides-pptx/slide-0001-figure-01.png").is_file())
        check = cli("validate_lecture.py", imported, "--original", slide_path,
                    "--output", self.validation, "--json")
        result = json.loads(check.stdout)
        self.assertEqual(result["original_validation"], "matching SHA256")
        self.assertEqual(result["source_sections"], 1)
        self.assertEqual(result["original_page_or_slide_count"], 1)
        self.assertFalse(result["blocking_issues"])

    def test_docx_paragraph_table_order_and_image(self):
        try:
            from docx import Document
            from docx.shared import Inches
            from PIL import Image
        except ImportError:
            self.skipTest("optional Word/Pillow unavailable")
        doc_path = self.root / "Anatomy handout.docx"
        picture = self.root / "tissue.png"
        png(picture)
        doc = Document()
        doc.add_paragraph("First, understand the renal cortex.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Structure"
        table.cell(0, 1).text = "Location"
        table.cell(1, 0).text = "Glomerulus"
        table.cell(1, 1).text = "Cortex"
        doc.add_paragraph("Then compare the medulla.")
        doc.add_picture(str(picture), width=Inches(1.0))
        doc.save(str(doc_path))
        cli("learn.py", doc_path, "--notes", self.notes, "--briefs", self.briefs)
        imported = self.notes / "anatomy-handout-docx.md"
        text = imported.read_text(encoding="utf-8")
        p1 = text.index("First, understand")
        tab = text.index("### Table 1")
        p2 = text.index("Then compare")
        self.assertTrue(p1 < tab < p2, "Tables must preserve their actual position")
        self.assertIn("Structure | Location", text)
        self.assertIn("DOCX figure; position not established", text)
        self.assertTrue(list((self.notes / "assets/anatomy-handout-docx").glob("doc-figure-*")))
        info = json.loads(cli("validate_lecture.py", imported, "--original", doc_path,
                              "--output", self.validation, "--json").stdout)
        self.assertEqual(info["original_validation"], "matching SHA256")
        self.assertIsNone(info["original_page_or_slide_count"])


if __name__ == "__main__":
    unittest.main()
